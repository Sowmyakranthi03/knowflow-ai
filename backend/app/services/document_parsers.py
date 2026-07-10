import logging
import re
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.config import settings
from app.core.exceptions import DocumentProcessingError
from app.schemas.document import ExtractedSection


logger = logging.getLogger(__name__)


class BaseDocumentParser(ABC):
    parser_name: str

    @abstractmethod
    def parse(self, file_path: Path) -> list[ExtractedSection]:
        """Extract structured sections from a document."""

    @staticmethod
    def normalize_text(text: str) -> str:
        normalized_lines: list[str] = []

        for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            cleaned_line = re.sub(r"[ \t]+", " ", line).strip()

            if cleaned_line:
                normalized_lines.append(cleaned_line)

        return "\n".join(normalized_lines).strip()

    @staticmethod
    def build_section(
        *,
        section_number: int,
        section_type: str,
        text: str,
        metadata: dict[str, Any] | None = None,
    ) -> ExtractedSection:
        normalized_text = BaseDocumentParser.normalize_text(text)

        return ExtractedSection(
            section_number=section_number,
            section_type=section_type,
            text=normalized_text,
            character_count=len(normalized_text),
            word_count=len(normalized_text.split()),
            metadata=metadata or {},
        )

    @staticmethod
    def validate_character_limit(
        sections: list[ExtractedSection],
    ) -> None:
        total_characters = sum(
            section.character_count
            for section in sections
        )

        if total_characters > settings.max_extracted_characters:
            raise DocumentProcessingError(
                "The extracted document text exceeds the processing limit.",
                status_code=413,
                error_code="EXTRACTED_TEXT_TOO_LARGE",
            )


class PdfDocumentParser(BaseDocumentParser):
    parser_name = "pypdf"

    def parse(self, file_path: Path) -> list[ExtractedSection]:
        try:
            reader = PdfReader(str(file_path))

            if reader.is_encrypted:
                try:
                    decryption_result = reader.decrypt("")
                except Exception as error:
                    raise DocumentProcessingError(
                        "Encrypted PDF documents are not supported.",
                        status_code=422,
                        error_code="ENCRYPTED_PDF",
                    ) from error

                if decryption_result == 0:
                    raise DocumentProcessingError(
                        "Encrypted PDF documents are not supported.",
                        status_code=422,
                        error_code="ENCRYPTED_PDF",
                    )

            sections: list[ExtractedSection] = []

            for page_index, page in enumerate(
                reader.pages,
                start=1,
            ):
                try:
                    extracted_text = page.extract_text() or ""
                except Exception as error:
                    logger.warning(
                        "Unable to extract PDF page %s from %s",
                        page_index,
                        file_path.name,
                    )

                    raise DocumentProcessingError(
                        f"Text could not be extracted from PDF page {page_index}.",
                        status_code=422,
                        error_code="PDF_PAGE_EXTRACTION_ERROR",
                    ) from error

                sections.append(
                    self.build_section(
                        section_number=page_index,
                        section_type="page",
                        text=extracted_text,
                        metadata={
                            "page_number": page_index,
                        },
                    )
                )

            if not sections:
                raise DocumentProcessingError(
                    "The PDF does not contain any pages.",
                    status_code=422,
                    error_code="EMPTY_PDF",
                )

            self.validate_character_limit(sections)

            return sections

        except DocumentProcessingError:
            raise

        except PdfReadError as error:
            raise DocumentProcessingError(
                "The PDF document could not be read.",
                status_code=422,
                error_code="PDF_READ_ERROR",
            ) from error

        except OSError as error:
            raise DocumentProcessingError(
                "The PDF file could not be accessed.",
                status_code=500,
                error_code="DOCUMENT_ACCESS_ERROR",
            ) from error


class DocxDocumentParser(BaseDocumentParser):
    parser_name = "python-docx"

    def parse(self, file_path: Path) -> list[ExtractedSection]:
        try:
            document = Document(str(file_path))
            sections: list[ExtractedSection] = []
            section_number = 1

            for paragraph_index, paragraph in enumerate(
                document.paragraphs,
                start=1,
            ):
                normalized_text = self.normalize_text(paragraph.text)

                if not normalized_text:
                    continue

                sections.append(
                    self.build_section(
                        section_number=section_number,
                        section_type="paragraph",
                        text=normalized_text,
                        metadata={
                            "paragraph_number": paragraph_index,
                            "style": (
                                paragraph.style.name
                                if paragraph.style is not None
                                else None
                            ),
                        },
                    )
                )

                section_number += 1

            for table_index, table in enumerate(
                document.tables,
                start=1,
            ):
                table_rows: list[str] = []

                for row in table.rows:
                    cells = [
                        self.normalize_text(cell.text)
                        for cell in row.cells
                    ]

                    table_rows.append(" | ".join(cells))

                table_text = "\n".join(table_rows)

                if not self.normalize_text(table_text):
                    continue

                sections.append(
                    self.build_section(
                        section_number=section_number,
                        section_type="table",
                        text=table_text,
                        metadata={
                            "table_number": table_index,
                            "row_count": len(table.rows),
                            "column_count": (
                                len(table.rows[0].cells)
                                if table.rows
                                else 0
                            ),
                        },
                    )
                )

                section_number += 1

            if not sections:
                raise DocumentProcessingError(
                    "The DOCX document does not contain extractable text.",
                    status_code=422,
                    error_code="NO_EXTRACTABLE_TEXT",
                )

            self.validate_character_limit(sections)

            return sections

        except DocumentProcessingError:
            raise

        except PackageNotFoundError as error:
            raise DocumentProcessingError(
                "The DOCX document could not be opened.",
                status_code=422,
                error_code="DOCX_READ_ERROR",
            ) from error

        except (ValueError, KeyError) as error:
            raise DocumentProcessingError(
                "The DOCX document structure is invalid.",
                status_code=422,
                error_code="DOCX_STRUCTURE_ERROR",
            ) from error

        except OSError as error:
            raise DocumentProcessingError(
                "The DOCX file could not be accessed.",
                status_code=500,
                error_code="DOCUMENT_ACCESS_ERROR",
            ) from error


class TxtDocumentParser(BaseDocumentParser):
    parser_name = "utf-8-text"

    def parse(self, file_path: Path) -> list[ExtractedSection]:
        try:
            raw_text = file_path.read_text(encoding="utf-8")
            normalized_text = self.normalize_text(raw_text)

            if not normalized_text:
                raise DocumentProcessingError(
                    "The TXT document does not contain extractable text.",
                    status_code=422,
                    error_code="NO_EXTRACTABLE_TEXT",
                )

            sections: list[ExtractedSection] = []

            paragraphs = re.split(
                r"\n\s*\n",
                raw_text.replace("\r\n", "\n").replace("\r", "\n"),
            )

            for paragraph_index, paragraph in enumerate(
                paragraphs,
                start=1,
            ):
                cleaned_paragraph = self.normalize_text(paragraph)

                if not cleaned_paragraph:
                    continue

                sections.append(
                    self.build_section(
                        section_number=len(sections) + 1,
                        section_type="paragraph",
                        text=cleaned_paragraph,
                        metadata={
                            "paragraph_number": paragraph_index,
                        },
                    )
                )

            if not sections:
                sections.append(
                    self.build_section(
                        section_number=1,
                        section_type="document",
                        text=normalized_text,
                    )
                )

            self.validate_character_limit(sections)

            return sections

        except DocumentProcessingError:
            raise

        except UnicodeDecodeError as error:
            raise DocumentProcessingError(
                "TXT documents must use UTF-8 encoding.",
                status_code=422,
                error_code="INVALID_TEXT_ENCODING",
            ) from error

        except OSError as error:
            raise DocumentProcessingError(
                "The TXT file could not be accessed.",
                status_code=500,
                error_code="DOCUMENT_ACCESS_ERROR",
            ) from error


class DocumentParserFactory:
    _parsers: dict[str, BaseDocumentParser] = {
        ".pdf": PdfDocumentParser(),
        ".docx": DocxDocumentParser(),
        ".txt": TxtDocumentParser(),
    }

    @classmethod
    def get_parser(
        cls,
        extension: str,
    ) -> BaseDocumentParser:
        parser = cls._parsers.get(extension.lower())

        if parser is None:
            raise DocumentProcessingError(
                "No parser is available for this document type.",
                status_code=415,
                error_code="PARSER_NOT_AVAILABLE",
            )

        return parser