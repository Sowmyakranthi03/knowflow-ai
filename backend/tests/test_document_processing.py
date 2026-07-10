import io
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from app.main import app
from app.repositories.document_repository import document_repository


client = TestClient(app)


@pytest.fixture
def processing_directories(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> tuple[Path, Path]:
    upload_directory = tmp_path / "uploads"
    processed_directory = tmp_path / "processed"

    upload_directory.mkdir()
    processed_directory.mkdir()

    monkeypatch.setattr(
        document_repository,
        "upload_directory",
        upload_directory,
    )

    monkeypatch.setattr(
        document_repository,
        "processed_directory",
        processed_directory,
    )

    return upload_directory, processed_directory


def create_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Employee Handbook", level=1)
    document.add_paragraph(
        "Employees must follow company security policies."
    )

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Department"
    table.cell(0, 1).text = "Contact"
    table.cell(1, 0).text = "IT"
    table.cell(1, 1).text = "support@example.com"

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def create_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(buffer)

    return buffer.getvalue()


def test_process_txt_document(
    processing_directories: tuple[Path, Path],
) -> None:
    upload_directory, processed_directory = processing_directories
    document_id = str(uuid4())

    document_path = upload_directory / f"{document_id}.txt"
    document_path.write_text(
        (
            "KnowFlow AI processes enterprise documents.\n\n"
            "It prepares text for semantic search."
        ),
        encoding="utf-8",
    )

    response = client.post(
        f"/api/v1/documents/{document_id}/process"
    )

    assert response.status_code == 200

    body = response.json()

    assert body["document_id"] == document_id
    assert body["file_extension"] == ".txt"
    assert body["parser"] == "utf-8-text"
    assert body["status"] == "processed"
    assert body["statistics"]["paragraph_count"] == 2
    assert body["statistics"]["word_count"] > 0
    assert len(body["sections"]) == 2

    processed_file = processed_directory / f"{document_id}.json"

    assert processed_file.exists()
    assert "KnowFlow AI" in processed_file.read_text(
        encoding="utf-8"
    )


def test_process_docx_document(
    processing_directories: tuple[Path, Path],
) -> None:
    upload_directory, processed_directory = processing_directories
    document_id = str(uuid4())

    document_path = upload_directory / f"{document_id}.docx"
    document_path.write_bytes(create_docx_bytes())

    response = client.post(
        f"/api/v1/documents/{document_id}/process"
    )

    assert response.status_code == 200

    body = response.json()

    assert body["parser"] == "python-docx"
    assert body["statistics"]["paragraph_count"] == 2
    assert body["statistics"]["table_count"] == 1
    assert "Employee Handbook" in body["full_text"]
    assert processed_directory.joinpath(
        f"{document_id}.json"
    ).exists()


def test_process_blank_pdf_returns_no_text_error(
    processing_directories: tuple[Path, Path],
) -> None:
    upload_directory, processed_directory = processing_directories
    document_id = str(uuid4())

    document_path = upload_directory / f"{document_id}.pdf"
    document_path.write_bytes(create_pdf_bytes())

    response = client.post(
        f"/api/v1/documents/{document_id}/process"
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == (
        "NO_EXTRACTABLE_TEXT"
    )
    assert list(processed_directory.iterdir()) == []


def test_process_missing_document(
    processing_directories: tuple[Path, Path],
) -> None:
    document_id = str(uuid4())

    response = client.post(
        f"/api/v1/documents/{document_id}/process"
    )

    assert response.status_code == 404
    assert response.json() == {
        "error": {
            "code": "DOCUMENT_NOT_FOUND",
            "message": (
                "The requested document could not be found."
            ),
        }
    }


def test_reject_invalid_document_id(
    processing_directories: tuple[Path, Path],
) -> None:
    response = client.post(
        "/api/v1/documents/../../secret/process"
    )

    assert response.status_code in {404, 405}


def test_process_empty_txt_document(
    processing_directories: tuple[Path, Path],
) -> None:
    upload_directory, processed_directory = processing_directories
    document_id = str(uuid4())

    document_path = upload_directory / f"{document_id}.txt"
    document_path.write_text(
        "   \n\n   ",
        encoding="utf-8",
    )

    response = client.post(
        f"/api/v1/documents/{document_id}/process"
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == (
        "NO_EXTRACTABLE_TEXT"
    )
    assert list(processed_directory.iterdir()) == []